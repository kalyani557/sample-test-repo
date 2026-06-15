import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

# 1. Generate a real PDF document
pdf_path = "security_policy.pdf"
c = canvas.Canvas(pdf_path, pagesize=letter)
c.setFont("Helvetica-Bold", 16)
c.drawString(100, 750, "GLOBAL DATA SECURITY STANDARD POLICY")
c.setFont("Helvetica-Bold", 12)
c.drawString(100, 710, "Section 1.1: Data Retention Requirements")
c.setFont("Helvetica", 10)
c.drawString(100, 690, "All user login metadata and application access logs must be securely stored")
c.drawString(100, 675, "for a minimum period of 3 years (36 months) from the initial creation date.")
c.setFont("Helvetica-Bold", 12)
c.drawString(100, 640, "Section 1.2: Cryptographic Controls")
c.setFont("Helvetica", 10)
c.drawString(100, 620, "All sensitive customer records must be encrypted at rest utilizing AES-256 protocols.")
c.save()
print(f"Created: {os.path.abspath(pdf_path)}")

# 2. Generate a real Word Document (.docx)
docx_path = "vendor_agreement.docx"
doc = Document()
doc.add_heading('ACME CORP VENDOR CONTRACT', level=0)
doc.add_heading('Termination Clauses', level=1)
doc.add_paragraph('Either party may terminate this agreement by providing at least 60 days written notice to the other party. Failure to provide adequate notice results in a standard $5,000 baseline penalty.')
doc.add_heading('Service Level Agreement (SLA)', level=1)
doc.add_paragraph('The vendor guarantees a system uptime SLA metric of 99.9%. Major support tickets must be acknowledged and answered within 4 business hours.')
doc.save(docx_path)
print(f"Created: {os.path.abspath(docx_path)}")